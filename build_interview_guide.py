from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_LINE_SPACING
from pathlib import Path

OUT = Path('LLM多模态视觉算法岗位面试题与参考答案.docx')

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_border(cell, color='D9D9D9', sz='6'):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tcPr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = 'w:' + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), sz)
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), color)

def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn('w:' + m))
        if node is None:
            node = OxmlElement('w:' + m)
            tcMar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement('w:tblHeader')
    tblHeader.set(qn('w:val'), 'true')
    trPr.append(tblHeader)

def set_font(run, name='Microsoft YaHei', size=10.5, bold=False, color='000000'):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:ascii'), name)
    run._element.rPr.rFonts.set(qn('w:hAnsi'), name)
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)

def add_para(doc, text='', style=None, bold_prefix=None, size=10.5, space_after=4, first_line=0):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.12
    if first_line:
        p.paragraph_format.first_line_indent = Inches(first_line)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix); set_font(r, size=size, bold=True)
        r2 = p.add_run(text[len(bold_prefix):]); set_font(r2, size=size)
    else:
        r = p.add_run(text); set_font(r, size=size)
    return p

sections = []
def section(title, items, intro=''):
    sections.append((title, intro, items))

section('一 Transformer 与 LLM 基础', [
('Transformer 整体结构是什么？', '典型 Transformer 由多层自注意力、前馈网络、残差连接和归一化组成。Encoder 处理双向上下文，Decoder 通过因果掩码逐 token 生成；GPT 类模型通常只保留 Decoder。'),
('Self-Attention 的计算过程是什么？', '输入先线性映射为 Q、K、V，��算注意力分数 S=QKᵀ/√d_k，经过因果掩码（如需要）和 Softmax 得到权重，再与 V 加权求和，最后经过输出投影。'),
('为什么 Attention 要除以 sqrt(d_k)？', '当 d_k 增大时，QKᵀ 的方差会随维度增大，Softmax 容易饱和、梯度变小。除以 √d_k 可以稳定分数分布，让训练更容易。'),
('MHA、MQA、GQA 有什么区别？', 'MHA 每个 Query 头都有独立 K/V；MQA 所有 Query 头共享一组 K/V；GQA 将 Query 头分组，每组共享一个 K/V 头，效果和显存开销介于两者之间。'),
('GQA 为什么能够降低 KV Cache 显存？', 'KV Cache 只缓存 K/V。GQA 减少 K/V 头数，缓存张量的头维度变小，所以每层、每个序列的缓存显存近似按 KV 头数同比下降。'),
('RoPE 的原理是什么？', 'RoPE 将每个二维特征对按位置角度旋转，使 Q 与 K 的内积天然包含相对位置信息，同时保留较好的长度外推特性。'),
('RMSNorm 和 LayerNorm 有什么区别？', 'LayerNorm 会减均值再按方差归一化；RMSNorm 只按均方根缩放，不做中心化，计算更简单，在大语言模型中通常能以更低开销获得相近稳定性。'),
('SwiGLU 为什么比普通 FFN 更常见？', 'SwiGLU 用 SiLU 门控另一支线性变换，能让网络按输入动态选择信息通路，通常带来更好的效果；代价是参数和计算略增，常通过调整中间维度控制。'),
('Pre-Norm 和 Post-Norm 有什么区别？', 'Pre-Norm 在子层前归一化，残差路径更直接，深层训练更稳定；Post-Norm 在子层后归一化，原始 Transformer 使用它，但深层或大模型训练往往更依赖精细的学习率和初始化。'),
('Decoder-only 为什么适合做 GPT 类模型？', '它用因果掩码建模“根据前文预测下一个 token”，训练目标简单、可扩展，推理时天然按序生成，适合统一处理对话、代码和文本生成任务。'),
('KV Cache 是什么？', '自回归生成时，历史 token 的 K、V 不随新 token 改变，可以缓存起来；下一步只计算新 token 的 Q/K/V，并复用历史 K/V。'),
('为什么 KV Cache 可以加速自回归生成？', '没有缓存时每一步都要重复计算整段前缀的 K/V；有缓存后每步只处理一个新 token，重复计算显著减少，代价是显存随上下文长度增长。'),
('KV Cache 的显存怎么计算？', '近似为层数×2（K、V）×序列长度×KV 头数×每头维度×batch×每元素字节数。GQA、低精度和分页管理都能降低实际占用。'),
('PagedAttention 解决了什么问题？', '它把 KV Cache 切成固定大小的页，由块表管理，避免连续大块显存和预留造成的碎片，支持请求动态增长、共享前缀和更高并发。'),
('为什么长上下文会导致显存压力？', 'KV Cache 与 token 数近似线性增长；同时注意力计算的访存和带宽压力增大，长序列还会降低可容纳的并发请求数。'),
('Prefill 和 Decode 有什么区别？', 'Prefill 一次处理输入提示词，计算密集并并行度高；Decode 每次生成一个 token，受显存带宽和 KV Cache 访问影响，更关注单步延迟与吞吐。'),
('TTFT、TPOT、TPS 分别是什么？', 'TTFT 是首 token 延迟；TPOT 是生成阶段平均每 token 延迟；TPS 是 tokens per second，通常等于生成 token 数除以生成耗时。三者要明确统计边界。'),
('LLM 推理过程中 CPU、GPU 分别承担什么工作？', 'GPU 负责矩阵乘、注意力和大部分模型算子；CPU 常负责请求调度、分词、采样、数据搬运、I/O 与部分后处理。端到端延迟还受 CPU-GPU 同步和内存拷贝影响。'),
])

section('二 模型训练与微调', [
('预训练、SFT、RLHF 分别解决什么问题？', '预训练从大规模无标注数据学习语言和世界知识；SFT 用高质量指令示例学习任务格式和行为；RLHF 用偏好信号优化回答风格、安全性与有用性。'),
('LoRA 为什么能够降低微调显存？', '冻结原模型权重，只训练低秩增量 ΔW=BA。优化器状态和梯度主要针对小矩阵，显著减少可训练参数和显存。'),
('LoRA 的 r、alpha 分别是什么？', 'r 是低秩矩阵的秩，决定适配容量；alpha 是缩放系数，实际增量常按 alpha/r 缩放。r 越大容量越强但参数和过拟合风险也更高。'),
('QLoRA 和 LoRA 有什么区别？', 'LoRA 通常在 FP16/BF16 基座上训练适配器；QLoRA 将冻结基座以 4-bit 量化存储，计算时反量化到低精度，并训练 LoRA，进一步降低显存。'),
('QLoRA 为什么可以在 4-bit 模型上训练？', '基座权重保持冻结，反向传播只更新 LoRA 参数；4-bit 权重在算子计算时临时反量化，因此不需要对量化权重本身保存完整梯度。'),
('NF4 和普通 INT4 有什么区别？', 'INT4 是均匀整数刻度；NF4 针对近似正态分布的权重设计非均匀码本，在相同位宽下通常有更小的量化误差。'),
('SFT 数据集应该怎么构造？', '覆盖真实任务分布，保证指令、上下文、答案格式清晰；去重去污染，过滤错误和敏感样本，并划分训练、验证集。对多轮数据要明确角色和 loss mask。'),
('SFT 训练时 loss 是怎么计算的？', '通常是自回归交叉熵，只对 assistant 的目标 token 计算 loss，用户指令和 padding 用 mask 忽略；token 级 loss 再按有效 token 聚合。'),
('为什么 instruction tuning 后模型能力可能下降？', '数据过窄、训练过久或学习率过大都会导致灾难性遗忘；格式化数据的分布也可能压制通用能力。可用混合通用数据、较小学习率、LoRA、早停和回归评测缓解。'),
('DPO 和 PPO 有什么区别？', 'PPO 需要策略模型、参考模型和奖励模型，通过在线采样并做裁剪更新；DPO 直接用偏好对的相对 log-prob 优化隐式奖励，流程简单、稳定且不必在线 rollout。'),
('DPO 为什么不需要显式训练 Reward Model？', '它利用最优策略与参考策略的关系，将奖励差写成策略 log-prob 比值，直接对 chosen/rejected 做分类式优化，相当于把奖励建模隐含在目标函数中。'),
('GRPO 和 PPO 有什么区别？', 'GRPO 对同一问题采样一组回答，用组内相对奖励估计优势，通常不单独训练 value model，因此更省显存；代价是依赖组采样质量和奖励设计。'),
('什么情况下应该选择 SFT，什么情况下应该选择 RAG？', '需要改变表达格式、行为模式或领域术语时优先 SFT；知识经常更新、需要可追溯来源或不适合写入参数时优先 RAG，二者也可组合。'),
('如果微调后模型出现灾难性遗忘，你怎么解决？', '先用通用能力回归集定位退化范围，再降低学习率或训练轮数，混入通用/旧域数据，减小 LoRA 容量，冻结更多层，并用正则或蒸馏约束输出。'),
])

section('三 Agent 与 LangGraph', [
('什么是 Agent？', 'Agent 是以模型为决策核心，能够感知状态、规划步骤、调用工具并根据结果迭代完成目标的系统。关键不在“会聊天”，而在闭环执行。'),
('Agent 和普通 LLM Workflow 有什么区别？', 'Workflow 的路径通常由代码预先确定；Agent 可根据当前状态动态决定下一步和工具。工程上常用固定流程包住可控的 Agent 节点。'),
('ReAct 是什么？', 'ReAct 把推理与行动交替进行：模型先形成下一步计划，再调用工具观察结果，基于观察继续决策，直到给出最终答案。'),
('Function Calling 是怎么工作的？', '开发者提供工具名称、参数 schema 和约束；模型输出结构化调用意图，运行时校验参数并执行函数，再把结果作为工具消息返回模型。'),
('Tool Calling 和 ReAct 有什么区别？', 'Tool Calling 是结构化调用机制；ReAct 是“思考—行动—观察”的控制范式。一个 ReAct Agent 可以���用 Tool Calling，二者不是互斥概念。'),
('LangGraph 为什么适合构建 Agent？', '它用有向图表达节点、条件边和循环，State 统一承载上下文，便于持久化、重试、人工介入、可观测和控制最大步数。'),
('StateGraph 中 State 是什么？', 'State 是图中节点共享的结构化状态，例如消息、任务阶段、工具结果、错误次数和用户权限。节点返回增量更新，由 reducer 合并。'),
('Agent 如何实现循环？', '让模型决策节点通过条件边回到工具节点或自身，工具结果再回到决策节点；当模型输出 final 或满足完成条件时走终止边。'),
('Agent 如何避免死循环？', '设置最大步数和总耗时，记录已执行动作与参数，检测重复调用，加入进度判定、失败预算和兜底终止；超过阈值时返回可解释的中间结果。'),
('Agent 最大执行轮数应该怎么设计？', '按任务类型和工具成本设置，短问答可为 3—6 轮，复杂任务可更高；应结合 token、时间、费用和成功率压测，而不是只拍一个固定数字。'),
('Tool 调用失败怎么办？', '区分可重试错误、参数错误和业务拒绝。可重试错误采用指数退避和次数上限；参数错误让模型修正 schema；不可恢复错误应尽快返回人工可处理的信息。'),
('Tool 返回错误后应该让 LLM 自己重试吗？', '可以，但必须把错误类型、可用参数和剩余重试次数结构化返回，并由策略限制重试；不能把无限重试责任交给模型。'),
('如何防止 Agent 越权调用 Tool？', '工具层做服务端鉴权和参数校验，按用户、角色、租户和资源范围授权；模型只拿到允许的工具列表，敏感操作增加审批、审计和幂等保护。'),
('多 Agent 为什么比单 Agent 更复杂？', '增加了角色边界、通信协议、共享状态、一致性、调度和成本问题；并行还会引入竞态与重复执行，必须有明确的协调者和终止条件。'),
('Multi-Agent 怎么进行任务分解？', '先定义目标与验收标准，再由 Supervisor 拆成互相独立、输入输出清晰的子任务；能并行的并行，存在依赖的按拓扑顺序串行。'),
('Supervisor Agent 是什么？', 'Supervisor 负责理解总目标、选择下一个专业 Agent、合并结果和决定是否结束；专业 Agent 只在自己的工具和职责范围内工作。'),
('多 Agent 如何共享状态？', '共享最小必要字段，采用带版本的状态存储或事件日志；对可变资源使用锁、幂等键或事务，避免各 Agent 直接覆盖彼此上下文。'),
('多 Agent 如何避免重复执行任务？', '为任务生成幂等键，持久化执行记录和状态机，执行前检查已完成/进行中状态；工具操作也要支持幂等和去重。'),
('为什么用 YAML + LangGraph，而不是把业务逻辑全部写死在 Python 里？', 'YAML 把业务规则与执行引擎解耦，便于非核心逻辑配置、审阅和版本管理；Parser 将 DSL 编译为 StateGraph，复杂控制仍保留在 Python 节点和 Tool 中，兼顾灵活性与可测试性。'),
])

section('四 Agent Memory', [
('Agent Memory 有哪些类型？', '常见有短期对话记忆、工作记忆、用户画像/偏好、情节记忆、语义知识记忆和工具/任务记忆。工程上应区分来源、时效、可信度和访问权限。'),
('Short-term Memory 和 Long-term Memory 有什么区别？', '短期记忆服务当前任务，通常在会话状态中，生命周期短、相关性高；长期记忆跨会话持久化，需检索、更新、过期和权限控制。'),
('Conversation History 应该全部放进 Prompt 吗？', '不应该。应按相关性、时间和 token 预算裁剪，配合摘要、最近窗口和检索记忆；敏感或低价值内容不应无条件注入。'),
('Memory 什么时间读取？', '在意图识别后、生成计划或回答前读取；先用当前查询构造���索条件，再按相关性、时效和权限过滤。'),
('Memory 什么时间写入？', '在任务完成或出现稳定事实后写入，避免把模型猜测和临时中间状态直接持久化。'),
('Memory 什么时间更新？', '当用户明确纠正、同一事实出现更高可信来源，或偏好发生变化时更新；保留版本和更新时间便于追溯。'),
('Memory 什么时间删除？', '用户要求删除、超过保留期、内容已过时或被判定为敏感/污染时删除或隔离，并同步清理向量索引和缓存。'),
('哪些 Memory 操作应该同步？', '会影响当前回答正确性的读取、权限判断和用户明确要求的更新应同步；写入失败不能静默改变当前任务语义。'),
('哪些 Memory 操作应该异步？', '摘要、统计、向量化、低优先级画像更新和过期清理可异步，使用队列、重试和最终一致性。'),
('如何压缩长期记忆？', '按主题合并重复事实，保留时间、来源和置信度；将多条情节摘要为可验证结论，定期做去重、过期和冲突检测。'),
('如何判断一条信息是否值得写入长期 Memory？', '看是否稳定、可复用、由用户确认或高可信来源支持，且未来检索能改变决策；一次性闲聊、猜测和敏感数据通常不写入。'),
('Memory 越多越好吗？', '不是。噪声会增加检索误召回、上下文冲突和隐私风险。目标是高相关、可解释、可更新的最小记忆集合。'),
('如何避免 Memory 污染？', '区分事实与推断，记录来源和置信度，写入前做 schema 校验，检索时做权限和冲突过滤，并支持用户查看、纠正和删除。'),
('如果用户修改了之前的信息，Memory 怎么更新？', '将新信息作为更高版本并标记旧事实失效；保留审计记录，更新结构化存储和向量索引，回答时只取当前有效版本。'),
('Agent 如何实现跨会话记忆？', '用稳定的用户/租户标识关联持久化存储，按查询检索相关记忆，再经权限和时效过滤后注入当前 State；避免把完整历史直接拼接。'),
])

section('五 RAG', [
('RAG 的完整流程是什么？', '文档解析与清洗→切块→Embedding/关键词索引→向量库与倒排库→查询改写或 HyDE→混合召回→去重与 Rerank→上下文组装→LLM 生成→引用与评估。'),
('为什么需要 Chunking？', '检索以 chunk 为粒度，合理切块能让主题集中、向量表达稳定，并控制上下文长度；应结合标题、段落、表格和语义边界切分。'),
('Chunk 太大有什么问题？', '主题混杂、向量被稀释、召回不精确，还会占用更多上下文 token，导致模型注意力被无关内容分散。'),
('Chunk 太小有什么问题？', '上下文不完整，实体、条件和结论容易被拆散，单块语义不足导致召回和回答都不稳定。'),
('Dense Retrieval 和 Sparse Retrieval 有什么区别？', 'Dense 用语义向量匹配同义表达；Sparse 依赖词项、BM25 等精确匹配，对专有名词、编号和罕见词更敏感，二者适合混合。'),
('BGE-M3 为什么适合混合检索？', '它支持多语言、长文本和多种检索模式，可提供 dense、sparse 等表示，便于用同一模型覆盖跨语言和关键词检索需求。'),
('HyDE 是什么？', 'HyDE 先让模型根据问题生成一个假设答案/文档，再对假设文本做 embedding 去检索真实文档，用生成语义桥接用户问题与知识库表达差异。'),
('RRF 是什么？', 'Reciprocal Rank Fusion 按排名倒数给多个召回列表加权，常用 score=Σ1/(k+rank)，无需把不同检索器分数强行校准。'),
('Reranker 为什么放在召回之后？', 'Cross-Encoder 逐对读取 query 与文档，精度高但成本大；先用便宜召回缩小候选，再 rerank 前几十条能平衡效果与延迟。'),
('Embedding Model 和 Reranker 有什么区别？', 'Embedding 将文本编码成可预计算向量用于大规模近似检索；Reranker 联合编码 query 与候选文档，输出更精细的相关性分数。'),
('RAG 为什么会召回错误文档？', '切块边界、查询表述、embedding 能力、索引质量、元数据过滤或领域术语都可能造成误召回；需看召回样本和检索指标定位。'),
('RAG 为什么会出现“召回正确但回答错误”？', '上下文排序不佳、证据不完整、模型忽略引用、提示词约束弱或模型本身推理错误。可增加证据拼接、引用约束、答案校验和拒答策略。'),
('如何评估 RAG？', '离线评估检索 Recall@K、MRR、NDCG 和答案 EM/F1/语义相似度；在线评估延迟、用户反馈、引用正确率、拒答率和端到端任务成功率。'),
('Recall、Precision、MRR、NDCG 分别是什么？', 'Recall 衡量相关项被找回的比例；Precision 衡量返回项中相关项比例；MRR 看第一个相关结果的倒数排名；NDCG 按位置和相关等级加权评估排序质量。'),
('如何处理长文档？', '先做层级解析，保留标题路径；采用父子 chunk、滑窗或摘要索引，召回小块后回填父段，并通过压缩和 rerank 控制上下文。'),
('如何处理 PDF 表格？', '优先做版面感知解析，保留行列关系、表头和单位；复杂表格转成结构化 Markdown/JSON，并把页码、表名和来源元数据一起索引。'),
('如何处理图片中的信息？', '先判断 OCR 是否足够；扫描文档可 OCR+版面解析，复杂图表或视觉关系用 VLM 描述，再把文本、坐标、页码和原图引用一并存储。'),
])

section('六 VLM 与多模态', [
('VLM 和普通 LLM 有什么区别？', 'VLM 在语言模型外增加视觉编码器和跨模态连接，让模型能把图像/视频转换成可与文本联合建模的表示。'),
('图片是怎么输入 LLM 的？', '图像先经 Vision Encoder 变成 patch 或区域特征，再由 projector 映射到 LLM 的词向量维度，作为视觉 token 与文本 token 一起进入 Transformer。'),
('Vision Encoder 的作用是什么？', '提取边缘、形状、物体、布局和语义等视觉特征，并将高维像素压缩为更适合语言模型处理的序列表示。'),
('CLIP 的基本思想是什么？', '用对比学习把图像编码和文本编码拉到同一向量空间，让匹配图文相似度高、不匹配组合相似度低，从而支持零样本分类和跨模态检索。'),
('图像 Token 是怎么进入 Transformer 的？', '图像特征经线性投影、位置编码或特殊标记处理后，与文本 embedding 拼接或通过 cross-attention 交互，注意力层因此能联合访问两种模态。'),
('VLM 为什么需要视觉投影层？', '视觉编码器和 LLM 的维度、分布与语义空间不同，projector 负责对齐维度和模态分布，是视觉信息进入语言空间的接口。'),
('图片分辨率提高为什么会增加计算量？', 'patch 数随分辨率平方增长，视觉 token 变多；后续注意力和投影的计算、显存及传输也随序列长度上升。'),
('多模态模型如何同时理解文字和图片？', '通过共享语言空间、交叉注意力或统一 token 序列建立对齐；训练中用图文对比、描述、问答和指令数据学习跨模态对应关系。'),
('OCR 和 VLM 有什么区别？', 'OCR 重点是识别文字及其位置，准确、可解释且适合批量文档；VLM 能理解场景、关系和图表，但文字精确率、成本和稳定性可能不如专用 OCR。'),
('什么场景适合 OCR，什么场景适合 VLM？', '票据、证件和表格字段抽取优先 OCR+版面模型；需要理解场景、图表趋势、物体关系或自然语言描述时使用 VLM，必要时组合。'),
('视频理解和图片理解有什么区别？', '视频增加时间维度，需要采样帧、建模运动和事件顺序；图片主要处理空间关系。视频还要面对更高 token、存储和延迟。'),
('视频为什么不能简单把所有帧全部送进 VLM？', '帧数会导致 token 和注意力成本快速增长，重复帧也会稀释关键信息。通常采用关键帧采样、镜头切分、低分辨率预筛和分层摘要。'),
])

section('七 传统视觉基础与 YOLO', [
('CNN 为什么适合处理图像？', '卷积具备局部感受野、权重共享和一定平移等变性，能以较少参数提取从边缘到语义的层级特征。'),
('卷积核到底在做什么？', '卷积核在局部窗口内对像素做加权求和并滑动，学习检测边缘、纹理、形状等模式，多个通道共同形成高层表示。'),
('Stride 是什么？', 'Stride 是卷积核每次滑动的步长。增大 stride 会下采样、减少特征图尺寸和计算，但可能损失细节。'),
('Padding 是什么？', 'Padding 在输入边缘补零或其他值，用于控制输出尺寸、保留边缘信息和让卷积层堆叠更稳定。'),
('Pooling 有什么作用？', '池化下采样并聚合局部响应，降低空间尺寸、计算量和小幅平移敏感性；现代网络也常用 stride convolution 替代。'),
('ResNet 为什么使用残差连接？', '残差连接让层学习相对输入的增量，提供更顺畅的梯度通路，缓解深层网络退化和梯度消失。'),
('BatchNorm 是什么？', 'BatchNorm 用 mini-batch 的均值和方差规范化激活，再用可学习的缩放��偏移恢复表达能力，可稳定训练并允许更大学习率。'),
('YOLO 的整体结构是什么？', 'YOLO 是单阶段检测器：Backbone 提取特征，Neck 做多尺度融合，Head 在各尺度输出类别、置信度和边界框。'),
('Backbone、Neck、Head 分别负责什么？', 'Backbone 负责通用视觉特征；Neck 融合不同分辨率、增强小目标信息；Head 将特征解码为分类、置信度和位置预测。'),
('YOLO 如何预测 Bounding Box？', '网络对网格/特征点回归中心、宽高或距离分布，并输出类别与 objectness；训练时用 IoU/GIoU/CIoU 等位置损失和分类损失。'),
('IoU 是什么？', 'IoU=预测框与真实框交集面积/并集面积，用来衡量定位重叠程度，也是正负样本、NMS 和评估的重要指标。'),
('NMS 是什么？', '按置信度排序，保留最高分框并删除与其 IoU 超过阈值的重叠框；类别感知 NMS 会按类别分别处理。'),
('为什么需要 NMS？', '同一目标常被多个网格或候选框重复预测，NMS 将重复结果合并为更少的最终检测框。'),
('Precision 和 Recall 是什么？', 'Precision=TP/(TP+FP)，表示预测为正的结果有多准；Recall=TP/(TP+FN)，表示真实目标被找回的比例，二者通常存在权衡。'),
('mAP 怎么计算？', '先对每个类别按置信度排序，画 Precision-Recall 曲线并求 AP，再对所有类别 AP 取平均得到 mAP；还要说明 IoU 阈值，如 mAP@0.5 或 mAP@0.5:0.95。'),
('AP 和 mAP 有什么区别？', 'AP 是单一类别在指定 IoU 条件下的平均精度；mAP 是多个类别 AP 的均值。'),
('YOLO 中正负样本怎么定义？', '根据预测点/候选框与真实框的匹配规则和 IoU、中心区域等条件分配正样本，其余为负样本或忽略样本；具体取决于 YOLO 版本。'),
('YOLO 推理速度为什么比较快？', '单阶段一次前向直接输出检测结果，避免候选区域逐个分类；卷积结构规则、可批量并行，配合多尺度特征和后处理可实现实时性。'),
])

section('八 目标跟踪与 ByteTrack', [
('什么是目标跟踪？', '在连续帧中为同一目标保持稳定 ID，并估计其位置、轨迹和状态。它通常以检测结果为观测，结合运动与外观进行数据关联。'),
('Detection 和 Tracking 有什么区别？', 'Detection 只回答当前帧有什么目标、位置和类别；Tracking 还要跨帧关联同一目标，处理遮挡、出入画面和 ID 连续性。'),
('Kalman Filter 在目标跟踪中做什么？', '根据运动模型预测下一帧状态，再用当前检测更新位置和速度，并通过协方差表示不确定性，为匹配提供预测框。'),
('Hungarian Algorithm 是什么？', '它解决最小代价二分图匹配，将轨迹与检测按 IoU 距离或外观代价进行全局一对一分配。'),
('ByteTrack 的核心思想是什么？', '先用高置信度检测框完成主要匹配，再将未匹配轨迹与低置信度检测框进行第二次匹配，从而恢复被遮挡或暂时低分的目标。'),
('ByteTrack 为什么还要利用低置信度检测框？', '低分不一定是假目标，可能是遮挡、模糊或姿态变化。若直接丢弃，会造成轨迹断裂和 ID Switch；结合运动预测可筛出有价值的低分框。'),
('DeepSORT 和 ByteTrack 有什么区别？', 'DeepSORT 依赖 ReID 外观特征与运动联合匹配，适合外观相似场景但计算更重；ByteTrack 主要依赖检测置信度和运动，简单快速、易部署。'),
('ReID 在 DeepSORT 中有什么作用？', '将目标裁剪成外观 embedding，在遮挡或交叉时辅助判断是否为同一目标，减少仅凭位置造成的错误关联。'),
('如果目标被遮挡了怎么办？', '使用 Kalman 预测保持轨迹，允许短时丢失；遮挡解除后先匹配高分框，再用低分框或外观特征恢复，超过最大丢失帧数才删除轨迹。'),
('ID Switch 是什么？', '同一真实目标的 ID 被错误地分配给另一个目标，或两个目标之间发生身份交换，是跟踪稳定性的关键错误。'),
('如何评价一个 Tracking 算法？', '检测层看 mAP，跟踪层常看 MOTA、IDF1、HOTA、ID Switch、轨迹完整率和 FPS；要同时报告精度、身份稳定性与实时性。'),
])

section('九 模型部署与优化', [
('PyTorch 模型如何部署到端侧？', '固定并验证输入输出→导出 ONNX 或目标格式→做算子与动态维度检查→选择 Execution Provider/加速后端→量化或融合→端侧联调→与原模型逐层和端到端验收。'),
('ONNX 是什么？', 'ONNX 是跨框架的计算图与算子标准格式，描述模型结构、权重和输入输出，使训练框架与推理引擎解耦。'),
('为什么需要 ONNX？', '便于在 PyTorch、TensorFlow、TensorRT、ONNX Runtime、RKNN 等工具之间迁移，并支持图优化、算子融合和跨硬件部署。'),
('ONNX Runtime 是什么？', 'ONNX Runtime 是通用推理引擎，提供图优化、内存规划和多种 Execution Provider，可在 CPU、CUDA、TensorRT、DirectML 等后端执行 ONNX。'),
('Execution Provider 是什么？', 'EP 是 ONNX Runtime 的硬件/后端适配层，负责把支持的算子分配到 CUDA、TensorRT、CPU 或其他设备，不同 EP 可按优先级回退。'),
('TensorRT 和 ONNX Runtime 有什么关系？', 'TensorRT 是 NVIDIA 的高性能推理优化器与运行时；ONNX Runtime 可通过 TensorRT EP 调用它，并负责图分区和不支持算子的回退。'),
('CUDA、TensorRT、ONNX Runtime 三者有什么关系？', 'CUDA 提供 GPU 编程与基础库；TensorRT 基于 CUDA 做图优化和 kernel 选择；ONNX Runtime 是统一推理接口，可调度 CUDA EP 或 TensorRT EP。'),
('FP32、FP16、INT8、INT4 有什么区别？', '它们分别用 32/16 位浮点和 8/4 位整数表示数值。位宽越低，显存和带宽通常越小，但动态范围和精度风险越大，需看硬件支持。'),
('量化为什么能够降低显存？', '权重和激活用更少 bit 存储与搬运，模型大小、KV 或中间缓冲区可下降，带宽压力也减小。'),
('量化为什么会损失精度？', '连续浮点值被映射到有限离散级别，异常值和敏感层产生更大误差；误差经过多层累积可能影响输出。'),
('PTQ 和 QAT 有什么区别？', 'PTQ 训练后直接用校准数据确定量化参数，成本低、上线快；QAT 在训练中模拟量化误差，让模型学会适应，通常精度更好但训练成本更高。'),
('INT8 量化具体怎么做？', '准备覆盖真实分布的 calibration dataset，统计权重/激活范围，选择 per-tensor 或 per-channel scale/zero-point，生成量化图并用代表集校准，最后做精度与延迟回归。'),
('Calibration 是什么？', 'Calibration 是用代表性样本估计激活分布和量化尺度的过程；数据覆盖不足会使端侧真实输入出现饱和或精度骤降。'),
('什么是算子融合？', '把连续算子合并为一个更大的计算单元，例如 Conv+BN+Activation，减少中间张量读写、kernel launch 和调度开销。'),
('什么是 Kernel Fusion？', '在 GPU kernel 层把多个小算子合并成一次 kernel 执行，减少访存和启动开销，常用于逐元素、归一化和注意力相关操作。'),
('模型推理为什么会出现 CPU/GPU 瓶颈？', '数据预处理、后处理、线程同步或拷贝可能卡在 CPU；算子不支持、工作集太小、访存受限或频繁 kernel launch 可能卡在 GPU。必须用 profiler 区分。'),
('如何降低模型端侧延迟？', '先 profile 定位热点，再考虑缩小输入、减少层或通道、FP16/INT8、算子融合、TensorRT/NNAPI 等后端、异步流水线与内存复用，并用相同口径复测。'),
('如何降低显存？', '降低权重和激活精度，减小 batch/输入和并发，启用内存复用与 KV 分页，裁剪模型或分层 offload；同时检查是否有重复拷贝和缓存未释放。'),
('如何提高吞吐量？', '通过动态 batching、请求合并、异步流水线、CUDA Graph、连续 batching、模型并行/多实例和更高效 kernel 提升单位时间处理量，同时监控尾延迟。'),
('如何定位模型推理慢在哪里？', '固定输入和预热次数，分别测预处理、H2D、推理、D2H、后处理和排队；结合 Nsight/ORT profiling 查看算子耗时、同步点、内存带宽与 kernel launch。'),
])

section('十 RKNN TensorRT NCNN', [
('TensorRT 的优势是什么？', '针对 NVIDIA GPU 做图优化、层融合、精度选择、kernel autotuning 和内存规划，通常能在稳定模型上获得低延迟和高吞吐。'),
('TensorRT FP16 和 INT8 有什么区别？', 'FP16 几乎不需校准、精度风险较小；INT8 需要 calibration 或 QAT，显存/带宽和延迟更优，但对敏感层和数据分布更敏感。'),
('TensorRT Engine 是什么？', 'Engine 是针对具体 GPU、输入形状和精度配置编译出的可序列化执行计划，包含层融合、kernel 选择和内存规划，换硬件或配置可能需要重新构建。'),
('TensorRT 如何进行 Layer Fusion？', '解析计算图后识别可合并的模式，如 Conv+BN、Conv+Activation，通过单 kernel 或插件执行，减少中间写回和调度。'),
('RKNN 是什么？', 'RKNN 是 Rockchip NPU 的模型转换、量化和推理工具链，常把 ONNX/TensorFlow 等模型转换为 RKNN 格式，在 RK3588 等芯片 NPU 上执行。'),
('RKNN 为什么适合 Rockchip 芯片？', '它针对 Rockchip NPU 的算子、内存和量化特性做了适配，能利用硬件加速并提供端侧 API；需要关注版本、支持算子和量化校准。'),
('NCNN 和 TensorRT 有什么区别？', 'NCNN 面向移动端和 CPU/Vulkan，依赖轻、跨平台、适合 Android/嵌入式；TensorRT 面向 NVIDIA GPU，优化和吞吐更强，但平台绑定更明显。'),
('ONNX → TensorRT 的基本流程是什么？', '导出并简化 ONNX→检查静态/动态 shape 和算子→配置 FP16/INT8→构建 engine→用代表输入校验输出→集成 runtime→在目标 GPU 上 benchmark。'),
('ONNX → RKNN 的基本流程是什么？', '准备合法 ONNX 和输入量化数据→用 RKNN toolkit 加载、配置 mean/std 与 target platform→build/quantize→导出 RKNN→在板端初始化 runtime 并校验精度、性能。'),
('如果 ONNX 模型存在 TensorRT 不支持的算子怎么办？', '先用 graph surgery 或改写模型替换等价算子；必要时实现 TensorRT plugin；也可让 ONNX Runtime/CPU 回退，但要评估数据拷贝和延迟影响。'),
('如果端侧模型精度突然下降，你怎么定位？', '先确认预处理、颜色通道、归一化、输入 layout 和后处理一致；再做 FP32→FP16→INT8 逐级对比，逐层 dump，检查算子替换、校准数据、动态 shape 和 NMS/阈值。'),
])

section('十一 性能优化场景题', [
('模型准确率可以，但是延迟 200ms，要求降到 50ms，你怎么办？', '先固定硬件、输入和测量口径并 profile，拆解预处理、推理、传输、后处理；再按收益尝试减小输入/模型、FP16、INT8、算子融合、TensorRT/ORT、异步流水线和内存复用，每一步都回归精度和 P50/P99 延迟。'),
('模型显存占用太高怎么办？', '按权重、激活、KV Cache、临时 workspace 和并发拆账；优先 FP16/INT8/INT4，降低 batch 和上下文，启用分页/内存复用，再考虑裁剪、共享权重和分层 offload。'),
('量化后模型精度下降很多，怎么办？', '检查 calibration 是否覆盖真实分布和预处理是否一致，定位异常值与敏感层，改用 per-channel/混合精度、保留关键层 FP16，必要时采用 QAT，并比较逐层误差与端到端指标。'),
])

section('十二 结合本地 AI 游戏助手项目', [
('为什么选择 GGUF？', 'GGUF 是 llama.cpp 生态常用的自包含模型格式，支持量化、元数据、分片和 mmap，便于跨平台本地加载，适合 Unity/端侧集成。'),
('为什么使用 Q4_K_M？', 'Q4_K_M 在模型大小、内存带宽和生成质量之间取得较好平衡；相比更激进的 3-bit 精度更稳，比 FP16 更省显存，适合消费级设备。'),
('4-bit 量化具体是什么？', '把权重从高精度数值映射到每个分组约 16 个离散级别，用 scale/zero-point 或码本恢复近似值；存储位宽约为 FP16 的四分之一，另有少量元数据开销。'),
('GGUF 和 ONNX 有什么区别？', 'GGUF 更偏生成式 LLM 的权重、tokenizer 和量化元数据封装，通常由 llama.cpp 执行；ONNX 是通用计算图交换格式，适合视觉模型和多种推理后端。'),
('LLMUnity 底层是什么？', '它通常负责 Unity 侧的模型加载、会话管理和流式输出封装，底层可调用 llama.cpp 或其他本地运行时；面试时应结合实际版本说明，不要把插件 API 和推理内核混为一谈。'),
('llama.cpp 为什么适合端侧？', 'C/C++ 实现、依赖少、支持 CPU/GPU 多后端、量化和 mmap，能在没有服务端的设备上以较低内存运行生成模型。'),
('如何测 TTFT？', '从提交请求或开始 prefill 的统一时间点计时，到收到第一个有效生成 token；需排除/单独报告模型加载、首次编译和网络等待，并在预热后重复取 P50/P95。'),
('30 tokens/s 是怎么测出来的？', '统计生成阶段实际输出 token 数，除以从首 token 到末 token 的耗时；明确是否包含采样、回调、UI 和停用 token，并报告平均值与分位数。'),
('TTS 的 RTF 是什么意思？', 'Real-time factor=音频处理耗时/音频时长。RTF<1 表示快于实时，例如生成 10 秒音频耗时 2 秒，RTF=0.2。'),
('为什么端到端延迟不是 TTFT + TTS？', 'LLM 可流式输出，TTS 可按句或按片段并行；首段文本到达后就能开始合成，端到端感知延迟取决于流水线重叠、缓冲和播放启动。'),
('如何实现 LLM 和 TTS 并行？', 'LLM 以句子边界或短片段流式输出，通过有界队列交给 TTS；TTS 合成与下一段 LLM 生成并行，队列需要背压、取消和顺序控制。'),
('CancellationToken 为什么可以降低打断延迟？', '取消信号能让生成、网络、TTS 和播放任务协作式退出，避免等待整段结果；应在模型回调、队列、I/O 和 UI 层传播并释放资源。'),
('2.5GB 显存是怎么控制的？', '对模型权重、KV Cache、激活、TTS 和图形渲染分别预算，选量化模型，限制上下文和并发，复用缓冲区，监控峰值而不是只看平均占用。'),
('为什么本地模型能够达到 600~800ms 的感知延迟？', '感知延迟主要由首 token、首段 TTS 和播放启动组成，而非完整回答时间；通过模型常驻、预热、量化、流式输出和 LLM/TTS 重叠可把可感知反馈提前。'),
])

section('十三 模拟面试官连环问', [
('你做的是 LLM，为什么可以胜任视觉/多模态算法岗位？', '我已有 Transformer、推理优化、数据与评测经验，能把这些能力迁移到 VLM 和视觉部署；同时我会用 YOLO、ByteTrack、CNN、量化和 ONNX/TensorRT 补齐视觉工程链路。我的优势是能把模型效果、Agent 逻辑和端侧落地一起考虑。'),
('Transformer 和 CNN 在视觉任务上的区别是什么？', 'CNN 的局部归纳偏置强，数据和计算效率好，适合稳定的局部纹理；Transformer 通过注意力建模长距离关系和全局上下文，但通常需要更多数据或预训练。实际工程中常采用 CNN、ViT 或混合结构。'),
('YOLO + ByteTrack 是怎么实现视频目标跟踪的？', '每帧先由 YOLO 输出类别、框和置信度；ByteTrack 用 Kalman 预测轨迹，先将高分检测与轨迹匹配，再用低分框补充匹配，更新轨迹、创建新 ID 并删除长期丢失轨迹。'),
('ByteTrack 为什么要利用低置信度检测框？', '遮挡��模糊会让真实目标分数下降，直接阈值过滤会让轨迹断裂。低分框先经过运动一致性匹配，能恢复目标，同时不把低分框直接当作新目标。'),
('ONNX Runtime 和 TensorRT 有什么区别？', 'ORT 是跨平台推理框架和统一 API，支持多种 EP、易部署；TensorRT 是 NVIDIA 专用优化器/运行时，通常在兼容模型上性能更高但平台绑定更强。ORT 也可通过 TensorRT EP 调用它。'),
('FP16、INT8、INT4 量化有什么区别？', 'FP16 是低精度浮点，兼容性和精度风险较小；INT8/INT4 是整数低 bit 表示，压缩和带宽收益更大，但需要校准、混合精度或 QAT 来控制误差。'),
('如果一个模型精度 95%，但是端侧延迟 200ms，要求 50ms，你怎么优化？', '先 profile 并确认 200ms 的组成，再按收益尝试输入和结构裁剪、FP16/INT8、算子融合、TensorRT/硬件 EP、异步流水线和内存复用；每步都以同一验证集检查精度，最后报告 P50/P95 和功耗。'),
('你的 Q4_K_M 是什么？为什么选择它？', '它是 GGUF/llama.cpp 中的一类 4-bit 分组量化方案，按不同层使用适合的块结构，通常在模型大小、带宽和质量之间平衡较好，适合本地实时交互。'),
('你的 Agent 如何防止死循环和 Tool 越权？', '图层设置最大步数、超时、重复调用检测和失败预算；工具层实施服务端鉴权、schema/资源校验、幂等与审计，敏感操作增加审批，模型只能看到被授权工具。'),
('如果让你把一个 PyTorch 的 YOLO 模型部署到 RK3588，你会怎么做？', '导出并简化 ONNX，固定或设计动态输入，检查 RKNN 支持算子；准备真实校准集，用 RKNN toolkit 转换和 INT8 量化，板端运行并逐层/端到端对齐，最后测 mAP、FPS、延迟、内存和功耗，必要时改写不支持算子或拆分 CPU/NPU。'),
])

def add_page_number(paragraph):
    run = paragraph.add_run()
    fldChar1 = OxmlElement('w:fldChar'); fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText'); instrText.set(qn('xml:space'), 'preserve'); instrText.text = 'PAGE'
    fldChar2 = OxmlElement('w:fldChar'); fldChar2.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar1); run._r.append(instrText); run._r.append(fldChar2)

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.7); sec.bottom_margin = Inches(0.65)
sec.left_margin = Inches(0.78); sec.right_margin = Inches(0.78)

styles = doc.styles
styles['Normal'].font.name = 'Microsoft YaHei'
styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
styles['Normal'].font.size = Pt(10.5)
for sname, size in [('Title', 24), ('Heading 1', 16), ('Heading 2', 12.5)]:
    st = styles[sname]
    st.font.name = 'Microsoft YaHei'
    st._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    st.font.color.rgb = RGBColor(0,0,0)
    st.font.size = Pt(size)
    st.font.bold = True

title = doc.add_paragraph(style='Title')
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
tr = title.add_run('LLM 多模态视觉算法岗位面试题与参考答案')
set_font(tr, size=24, bold=True)
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.paragraph_format.space_after = Pt(18)
sr = sub.add_run('覆盖 LLM、Agent、RAG、VLM、视觉算法、目标跟踪与端侧部署优化')
set_font(sr, size=11, color='4F5B66')

add_para(doc, '使用说明', style='Heading 1', size=16, space_after=6)
add_para(doc, '这份资料把岗位准备重点整理成一套可直接口述的面试答案。回答优先给出结论，再补充原理、工程取舍和验证方法。面试时不必逐字背诵，应结合自己简历中的项目、指标和踩坑经历，形成“原理—实现—结果—反思”的表达。', size=10.5, space_after=10)

add_para(doc, '准备优先级', style='Heading 2', size=12.5, space_after=5)
tbl = doc.add_table(rows=1, cols=3)
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
tbl.autofit = False
widths = [Inches(0.7), Inches(2.1), Inches(3.9)]
headers = ['顺序', '模块', '建议重点']
for i, h in enumerate(headers):
    c = tbl.rows[0].cells[i]; c.width = widths[i]; set_cell_shading(c, '1F4E79'); set_cell_border(c); set_cell_margins(c)
    c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(h); set_font(r, size=10, bold=True, color='FFFFFF')
set_repeat_table_header(tbl.rows[0])
priority_rows = [
    ('1', '端侧部署与性能', 'ONNX、TensorRT、RKNN、量化、Profiling、延迟与显存优化'),
    ('2', '视觉基础与跟踪', 'CNN、YOLO、IoU、NMS、mAP、Kalman、Hungarian、ByteTrack'),
    ('3', 'Transformer 与推理', '注意力、RoPE、KV Cache、Prefill/Decode、服务指标'),
    ('4', 'Agent 与 RAG', 'LangGraph、Tool Calling、Memory、混合检索、评估'),
    ('5', 'VLM 与项目表达', '视觉输入链路、OCR/VLM 选择、GGUF、流式 LLM+TTS'),
]
for ridx, row in enumerate(priority_rows):
    cells = tbl.add_row().cells
    for i, val in enumerate(row):
        cells[i].width = widths[i]; set_cell_border(cells[i]); set_cell_margins(cells[i]); cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if ridx % 2 == 1: set_cell_shading(cells[i], 'F3F6F9')
        p = cells[i].paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(val); set_font(r, size=9.5)

add_para(doc, '目录', style='Heading 1', size=16, space_after=6)
for idx, (title_s, intro, items) in enumerate(sections, start=1):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2)
    r = p.add_run(f'{idx:02d}  {title_s}（{len(items)} 题）'); set_font(r, size=10.5, bold=True)
add_para(doc, '附录：答题方法与复习清单', style='Heading 2', size=12.5, space_after=8)
add_para(doc, '建议复习节奏：第一轮通读并标记不会的术语；第二轮用 60—90 秒口述每题；第三轮把答案替换成自己的项目证据。对于部署和视觉题，最好准备一张从模型导出到端侧验收的流程图，以及一组延迟、显存、精度和功耗数据。', size=10.5, space_after=12)

for idx, (title_s, intro, items) in enumerate(sections, start=1):
    if idx > 1:
        doc.add_page_break()
    p = doc.add_paragraph(style='Heading 1'); p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(6)
    r = p.add_run(f'{idx:02d}  {title_s}'); set_font(r, size=16, bold=True)
    if intro:
        add_para(doc, intro, size=10.5, space_after=8)
    for local_idx, (q, a) in enumerate(items, start=1):
        p = doc.add_paragraph(style='Heading 2')
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(5 if local_idx > 1 else 2)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(f'{q}'); set_font(r, size=11.5, bold=True)
        add_para(doc, '参考答案：' + a, bold_prefix='参考答案：', size=10.2, space_after=5)

add_para(doc, '附录 答题方法与复习清单', style='Heading 1', size=16, space_after=6)
for text in [
    '通用结构：先给一句结论，再讲关键机制，最后补充工程取舍、失败边界和如何验证。',
    '项目深挖：准备项目背景、个人负责部分、技术选型、关键指标、一次失败排查和下一步改进。',
    '部署题：任何“怎么优化”都先回答测量口径与 Profiling，再按模型、精度、算子、运行时、流水线五层推进。',
    '视觉题：把 YOLO 检测、ByteTrack 关联、Kalman 预测、Hungarian 匹配和 IDF1/HOTA 评价串成一条链。',
    'Agent 题：把模型决策和工具执行分开，强调 State、权限、幂等、超时、重试、审计与可观测性。',
]:
    p = doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text); set_font(r, size=10.5)

for section_obj in doc.sections:
    footer = section_obj.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('LLM 多模态视觉算法岗位面试准备  ·  '); set_font(r, size=8.5, color='6B7280')
    add_page_number(p)

doc.core_properties.title = 'LLM 多模态视觉算法岗位面试题与参考答案'
doc.core_properties.subject = 'LLM、Agent、RAG、VLM、视觉算法、目标跟踪与模型部署面试准备'
doc.core_properties.author = 'Codex'
doc.save(OUT)
print(OUT.resolve())
